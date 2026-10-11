#!/usr/bin/env python3
"""Lema humanizer runner. The agent using this skill does the LLM steps; this script does the rest.

Agent mode (default, no API keys - the agent is the language model):
  prepare FILE        split FILE into prose blocks and create <name>.lema-work/ with
                      originals.md plus zh.md and ja.md templates for the agent to fill
  nmt WORKDIR         Google Translate JA > FI > EN on ja.md, heuristic cleanup and integrity
                      checks; writes en.md with a clean/FLAGGED note per block
  assemble WORKDIR    re-check en.md and write <name>.lema.<ext> next to FILE

API mode (optional; needs the Lema engine repo and its provider keys in <Lema repo>/.env):
  api FILE [--engine v2|v3|v1 ...]      api --text "..."

Other:
  preview FILE              list the paragraphs that would be rewritten
  verify OUTPUT ORIGINAL    re-check an edited output against the original
  check-setup               bundled pipeline, Google Translate reachability, optional api mode

Only prose paragraphs are rewritten. Headings, tables, lists, quotes, code, captions, short
fragments and everything after a References heading are left exactly as they are.
"""
from __future__ import annotations

import argparse
import contextlib
import json
import os
import re
import shutil
import sys
import threading
import time
import unicodedata
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

PROVIDER_KEYS = {
    "Bedrock": "AWS_BEDROCK_API_KEY",
    "Groq": "GROQ_API_KEY",
    "NVIDIA": "NVIDIA_API_KEY",
    "Cerebras": "CEREBRAS_API_KEY",
}
PROVIDER_ORDER = ["Bedrock", "Groq", "NVIDIA", "Cerebras"]

REF_HEADINGS = {
    "references", "reference list", "bibliography", "works cited",
    "sources", "literature cited", "citations",
}
HEADING_RE = re.compile(r"^\s{0,3}#{1,6}(\s|$)")
SETEXT_RE = re.compile(r"^\s{0,3}(=+|-+)\s*$")
LIST_RE = re.compile(r"^\s{0,3}(?:[-*+•]|\d{1,3}[.)]|\(?[a-zA-Z0-9]{1,3}\))\s+")
HR_RE = re.compile(r"^\s{0,3}(?:-{3,}|\*{3,}|_{3,})\s*$")
FENCE_RE = re.compile(r"^\s{0,3}(```|~~~)")
CAPTION_RE = re.compile(r"^(?:table|figure|fig\.?|chart|exhibit)\s*[\dIVX]", re.I)
LINK_RE = re.compile(r"https?://")
XML_BAD_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")

YEAR = r"(?:1[6-9]|20)\d{2}[a-z]?"
YEAR_RE = re.compile(r"\b" + YEAR + r"\b")
NUM_RE = re.compile(r"(?<![\w.])\d+(?:[.,]\d+)*%?")
ACRONYM_RE = re.compile(r"\b[A-Z][A-Z0-9]*[A-Z0-9](?:[-–/][A-Z0-9]+)*\b")
NAME = r"[A-Z][\w'’\-]+"
NARRATIVE_CITE_RE = re.compile(
    r"(" + NAME + r"(?:\s+(?:and|&)\s+" + NAME + r")?(?:\s+et\s+al\.?)?)(?:['’]s)?"
    r"\s+\((" + YEAR + r"(?:,\s*" + YEAR + r")*)\)"
)
PAREN_CITE_RE = re.compile(r"\(([^()]*?\b" + YEAR + r"\b[^()]*)\)")
AUTHOR_SPLIT_RE = re.compile(r"\s*(?:,|&|\band\b|\bet\s+al\.?)\s*")
CAP_WORD = r"[A-Z][\w'’\-]*"
PROPER_PHRASE_RE = re.compile(CAP_WORD + r"(?:\s+(?:of|for|and|the|de|on)?\s*" + CAP_WORD + r")+")
PHRASE_STOPWORDS = {
    "the", "a", "an", "this", "that", "these", "those", "their", "its", "our", "in", "on", "at", "for",
    "when", "where", "while", "both", "neither", "each", "such", "by", "with", "from", "as", "after",
}
# Letters that English prose almost never uses: Finnish/Swedish vowels, CJK, kana, Hangul, Cyrillic.
FOREIGN_CHAR_RE = re.compile(r"[äöåÄÖÅЀ-ӿ぀-ヿ㐀-鿿가-힯]")
FINNISH_WORDS = {"ja", "että", "sekä", "jotka", "joka", "tämä", "mutta", "myös",
                 "ovat", "eivät", "kuin", "tai"}
FIRST_PERSON_RE = re.compile(r"\b(?:we|our|ours|ourselves)\b", re.I)
KANA_RE = re.compile(r"[぀-ヿ]")
HAN_RE = re.compile(r"[一-鿿]")

MARKER_RE = re.compile(r"<!--\s*block\s+(\d+)\b.*?-->", re.S)
PLACEHOLDER = "TODO"
# Zero-width and soft-hyphen characters Google Translate sometimes leaves in its output.
INVISIBLE_RE = re.compile(r"[­​‌‍⁠﻿]")
# Lema's heuristic_cleanup() swaps "Therefore" for a lowercase "so" (and similar) at sentence starts.
LOWERED_START_RE = re.compile(r"(^|[.!?]\s+)(so|also|as a result|in the end)\b")

_print_lock = threading.Lock()


def log(msg: str) -> None:
    with _print_lock:
        print(msg, file=sys.__stderr__, flush=True)


def tidy_output(text: str) -> str:
    """Drop invisible characters, turn the non-breaking hyphens gpt-oss emits into plain ones, and
    re-capitalise sentence starts that Lema's cleanup lowercased."""
    text = INVISIBLE_RE.sub("", text).replace("‑", "-").replace("‐", "-")
    text = LOWERED_START_RE.sub(lambda m: m.group(1) + m.group(2)[0].upper() + m.group(2)[1:], text)
    text = re.sub(r"[ \t]+([,.;:])", r"\1", text)  # "instrument , but" left by dash stripping
    return re.sub(r"[ \t]{2,}", " ", text).strip()


# ── The bundled pipeline, and the optional Lema engine repo for api mode ──────

def import_pipeline():
    """The bundled keyless half of the v2 pipeline: Google Translate hops + heuristic cleanup."""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    warnings.filterwarnings("ignore")
    try:
        import pipeline
    except ImportError as e:
        sys.exit(f"Could not import pipeline.py from next to this script: {e}. "
                 "Run: pip install requests deep-translator")
    return pipeline


def find_lema_home(explicit: str | None) -> Path:
    """api mode only: the Lema engine repo that holds v2/v3/v1 and their provider calls."""
    here = Path(__file__).resolve()
    candidates = [explicit, os.environ.get("LEMA_HOME")]
    pointer = here.parent.parent / "lema_home.txt"  # written by install.py --lema-home
    if pointer.exists():
        candidates.append(pointer.read_text(encoding="utf-8").strip())
    for c in candidates:
        if c and (Path(c) / "v2" / "v2_engine.py").exists():
            return Path(c)
    sys.exit(
        "api mode needs the Lema engine repo (looked for v2/v2_engine.py), which is not part of "
        "this skill. Pass --lema-home <path>, set LEMA_HOME, or use the default agent mode "
        "(prepare / nmt / assemble), which needs no engine repo and no API keys."
    )


def load_keys(lema_home: Path, env_files: list[str]) -> list[str]:
    try:
        from dotenv import load_dotenv
    except ImportError:
        sys.exit("python-dotenv is missing. Run: pip install -r \"%s\"" % (lema_home / "requirements.txt"))
    # Real environment variables win, then --env-file, then the repo's own .env files.
    for f in env_files:
        if not Path(f).exists():
            sys.exit(f"--env-file not found: {f}")
        load_dotenv(f, override=False)
    for f in (lema_home / ".env", lema_home / "backend" / ".env"):
        if f.exists():
            load_dotenv(f, override=False)
    return [p for p in PROVIDER_ORDER if os.environ.get(PROVIDER_KEYS[p])]


def provider_chain(preferred: str, available: list[str], fallback: bool) -> list[str]:
    if preferred not in available:
        sys.exit(
            f"No key for {preferred}: set {PROVIDER_KEYS[preferred]} in the environment or in "
            f"<LEMA_HOME>/.env. Providers with keys: {', '.join(available) or 'none'}."
        )
    if not fallback:
        return [preferred]
    return [preferred] + [p for p in available if p != preferred]


def load_engine(name: str, lema_home: Path, mode: str, intensity: str, strip_dashes: bool):
    sys.path.insert(0, str(lema_home))
    warnings.filterwarnings("ignore")
    try:
        if name == "v2":
            from v2.v2_engine import humanize_v2 as fn
        elif name == "v3":
            from v3.v3_engine import humanize_v3 as fn
        else:
            from humanizer import process_text
    except ImportError as e:
        sys.exit(f"Could not import Lema engine {name}: {e}. Run: pip install -r \"{lema_home / 'requirements.txt'}\"")

    if name in ("v2", "v3"):
        def run(text: str, provider: str, model: str | None) -> str:
            return fn(text=text, provider=provider, api_key=None, model=model)["humanized"]
    else:
        def run(text: str, provider: str, model: str | None) -> str:
            out, _temps = process_text(
                input_text=text, api_key="", mode=mode, intensity=intensity,
                provider=provider, model=model, use_sabotage=False, remove_ehm=strip_dashes,
            )
            return out
    return run


# ── Integrity checks ──────────────────────────────────────────────────────────

def fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.replace("’", "'")
    for dash in ("‐", "‑", "–", "-"):
        s = s.replace(dash, " ")
    return re.sub(r"\s+", " ", s).lower()


def squash(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("’", "'")).strip()


def capitalised_mid_sentence(word: str, text: str) -> bool:
    for m in re.finditer(r"\b" + re.escape(word) + r"\b", text):
        before = text[: m.start()].rstrip()
        if before and before[-1] not in ".!?:":
            return True
    return False


def protected_items(text: str) -> dict:
    names: set[str] = set()
    exact: set[str] = set()
    for m in NARRATIVE_CITE_RE.finditer(text):
        exact.add(squash(m.group(0)))
        for part in AUTHOR_SPLIT_RE.split(m.group(1)):
            if part[:1].isupper():
                names.add(re.sub(r"['’]s$", "", part))
    for m in PAREN_CITE_RE.finditer(text):
        found_cite = False
        for chunk in m.group(1).split(";"):
            cm = re.match(r"\s*(?:e\.g\.,?\s*|see\s+|cf\.\s*)?(.+?),\s*" + YEAR, chunk)
            if not cm:
                continue
            authors = [a for a in AUTHOR_SPLIT_RE.split(cm.group(1)) if a[:1].isupper()]
            if authors:
                found_cite = True
                names.update(authors)
        if found_cite:
            exact.add(squash(m.group(0)))
    phrases: set[str] = set()
    for m in PROPER_PHRASE_RE.finditer(text):
        words = m.group(0).split()
        sentence_start = not text[: m.start()].strip() or text[: m.start()].rstrip()[-1] in ".!?:"
        if sentence_start and words and not capitalised_mid_sentence(words[0], text):
            words = words[1:]  # "Before H3 ..." - capitalised only because it starts a sentence
        while words and (words[0].lower() in PHRASE_STOPWORDS or re.search(r"['’]s$", words[0])):
            words = words[1:]
        while words and words[-1].lower() in {"of", "for", "and", "the", "de", "on"}:
            words = words[:-1]
        if words:
            words[-1] = re.sub(r"['’]s$", "", words[-1])
        caps = [w for w in words if w[:1].isupper()]
        if len(caps) >= 2 and not all(w.isupper() or w in names for w in caps):
            phrases.add(" ".join(words))
    return {
        "names": names,
        "numbers": set(NUM_RE.findall(text)) | set(YEAR_RE.findall(text)),
        "acronyms": set(ACRONYM_RE.findall(text)),
        "phrases": phrases,
        "citations": exact,
    }


def lost_items(items: dict, text: str) -> list[str]:
    """Protected items from the original that `text` no longer contains (works in mixed-script text)."""
    folded = fold(text)
    lost = [f"author '{n}'" for n in sorted(items["names"]) if fold(n) not in folded]
    lost += [f"number '{n}'" for n in sorted(items["numbers"])
             if not re.search(r"(?<![\d.,])" + re.escape(n) + r"(?!\d)", text)]
    lost += [f"acronym '{a}'" for a in sorted(items["acronyms"])
             if not re.search(r"(?<![A-Za-z0-9])" + re.escape(a) + r"(?![A-Za-z0-9])", text)]
    lost += [f"name/term '{p}'" for p in sorted(items["phrases"]) if fold(p) not in folded]
    return lost


def untranslated_words(original: str, rewritten: str) -> list[str]:
    before = set(re.findall(r"\S+", original))
    odd = [w for w in re.findall(r"\S+", rewritten) if FOREIGN_CHAR_RE.search(w) and w not in before]
    fi_hits = [w for w in re.findall(r"\b\w+\b", rewritten.lower()) if w in FINNISH_WORDS]
    if len(fi_hits) >= 2:
        odd += fi_hits
    return odd


def check_integrity(original: str, rewritten: str) -> tuple[list[str], list[str]]:
    """Return (hard, soft) problems. Hard problems trigger a retry in api mode."""
    items = protected_items(original)
    hard: list[str] = []
    soft: list[str] = []
    foreign = untranslated_words(original, rewritten)
    if foreign:
        hard.append("untranslated text: '" + " ".join(foreign[:5]) + "'")
    hard += lost_items(items, rewritten)
    out_squash = squash(rewritten)
    soft += [f"citation wording '{c}'" for c in sorted(items["citations"]) if c not in out_squash]
    if FIRST_PERSON_RE.search(rewritten) and not FIRST_PERSON_RE.search(original):
        soft.append("switched to first person ('we/our')")
    words_in, words_out = len(original.split()), len(rewritten.split())
    if words_in:
        ratio = words_out / words_in
        if ratio < 0.65:
            hard.append(f"length dropped to {ratio:.0%} of original (possible omission)")
        elif ratio > 1.6:
            soft.append(f"length grew to {ratio:.0%} of original")
    return hard, soft


# ── Splitting documents into blocks ───────────────────────────────────────────

def strip_inline_md(s: str) -> str:
    s = re.sub(r"(\*\*|__)(?=\S)(.+?)(?<=\S)\1", r"\2", s)
    s = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"\1", s)
    s = re.sub(r"(?<![\w_])_(?=\S)(.+?)(?<=\S)_(?![\w_])", r"\1", s)
    return s


def ends_with_sentence_punct(s: str) -> bool:
    s = s.strip().rstrip("\"'”’)*_")
    return bool(s) and s[-1] in ".?!:;"


def heading_text(line: str) -> str:
    t = re.sub(r"^\s{0,3}#{1,6}\s*", "", line)
    t = re.sub(r"[*_#`]", "", t)
    t = re.sub(r"^[\dIVX]+[.)]?\s+", "", t.strip())
    return t.strip().rstrip(".:").lower()


def split_markdown(text: str) -> list[dict]:
    lines = text.split("\n")
    segs: list[dict] = []
    i, n = 0, len(lines)
    if lines and lines[0].strip() == "---":
        for j in range(1, n):
            if lines[j].strip() in ("---", "..."):
                segs.append({"kind": "frontmatter", "lines": lines[: j + 1], "line": 1})
                i = j + 1
                break
    while i < n:
        line = lines[i]
        if not line.strip():
            segs.append({"kind": "blank", "lines": [line], "line": i + 1})
            i += 1
            continue
        m = FENCE_RE.match(line)
        if m:
            j = i + 1
            while j < n and not lines[j].lstrip().startswith(m.group(1)):
                j += 1
            segs.append({"kind": "code", "lines": lines[i : j + 1], "line": i + 1})
            i = j + 1
            continue
        if HEADING_RE.match(line):
            segs.append({"kind": "block", "lines": [line], "line": i + 1})
            i += 1
            continue
        j = i + 1
        while j < n and lines[j].strip() and not FENCE_RE.match(lines[j]) and not HEADING_RE.match(lines[j]):
            j += 1
        segs.append({"kind": "block", "lines": lines[i:j], "line": i + 1})
        i = j
    return segs


def classify_markdown(segs: list[dict]) -> None:
    in_refs = False
    for seg in segs:
        if seg["kind"] != "block":
            continue
        lines = seg["lines"]
        first = lines[0].lstrip()
        joined = " ".join(l.strip() for l in lines)
        plain = strip_inline_md(joined)
        is_heading = bool(HEADING_RE.match(lines[0])) or (len(lines) == 2 and SETEXT_RE.match(lines[1]))
        if is_heading:
            in_refs = in_refs or heading_text(lines[0]) in REF_HEADINGS
            seg["kind"] = "heading"
        elif in_refs:
            seg["kind"] = "references"
        elif any(l.lstrip().startswith("|") for l in lines):
            seg["kind"] = "table"
        elif LIST_RE.match(lines[0]):
            seg["kind"] = "list"
        elif first.startswith(">"):
            seg["kind"] = "quote"
        elif first.startswith("<"):
            seg["kind"] = "html"
        elif len(lines) == 1 and HR_RE.match(lines[0]):
            seg["kind"] = "rule"
        elif plain.strip().rstrip(".:").lower() in REF_HEADINGS:
            in_refs = True
            seg["kind"] = "heading"
        elif LINK_RE.search(joined):
            seg["kind"] = "has-link"
        elif re.fullmatch(r"(\*\*|__|\*|_).+\1", joined.strip()):
            seg["kind"] = "emphasised-line"
        elif CAPTION_RE.match(plain) or plain.lower().startswith(("keywords", "key words")):
            seg["kind"] = "caption"
        elif len(plain.split()) < 5 or (not ends_with_sentence_punct(plain) and len(plain.split()) < 12):
            seg["kind"] = "fragment"
        else:
            seg["kind"] = "prose"
            seg["text"] = plain
            seg["had_formatting"] = plain != joined


def docx_targets(doc) -> list[dict]:
    targets = []
    in_refs = False
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        low = t.lower().rstrip(".:")
        style = (p.style.name if p.style is not None else "").lower()
        kind = "prose"
        if low in REF_HEADINGS:
            in_refs, kind = True, "heading"
        elif in_refs:
            kind = "references"
        elif any(h in style for h in ("heading", "title", "subtitle", "caption", "header", "footer", "toc")):
            kind = "heading"
        elif any(h in style for h in ("list", "bullet", "number")) or LIST_RE.match(t):
            kind = "list"
        elif CAPTION_RE.match(t) or low.startswith(("keywords", "key words")):
            kind = "caption"
        elif getattr(p, "hyperlinks", None) or LINK_RE.search(t):
            kind = "has-link"
        elif len(t.split()) < 5 or (not ends_with_sentence_punct(t) and len(t.split()) < 12):
            kind = "fragment"
        targets.append({"kind": kind, "para": i, "line": i + 1, "text": t})
    return targets


def load_source(src: Path) -> dict:
    """Parse a .md/.txt/.docx file. Prose blocks get ids 1..N in document order."""
    if not src.exists():
        sys.exit(f"Input not found: {src}")
    if src.suffix.lower() == ".docx":
        import docx
        doc = docx.Document(str(src))
        items = docx_targets(doc)
        blocks = [t for t in items if t["kind"] == "prose"]
        kept, kinds_source = len(items) - len(blocks), items
        container = {"kind": "docx", "doc": doc}
    else:
        raw = src.read_bytes().decode("utf-8-sig")
        segs = split_markdown(raw.replace("\r\n", "\n"))
        classify_markdown(segs)
        blocks = [s for s in segs if s["kind"] == "prose"]
        kinds_source = [s for s in segs if s["kind"] != "blank"]
        kept = len(kinds_source) - len(blocks)
        container = {"kind": "text", "segs": segs, "newline": "\r\n" if "\r\n" in raw else "\n"}
    for i, b in enumerate(blocks, 1):
        b["id"] = i
    kinds: dict[str, int] = {}
    for s in kinds_source:
        kinds[s["kind"]] = kinds.get(s["kind"], 0) + 1
    return {"container": container, "blocks": blocks, "kept": kept, "kinds": kinds}


def write_document(container: dict, out: Path, blocks: list[dict]) -> None:
    usable = {b["id"] for b in blocks if b.get("result") and b.get("status") != "error"}
    if container["kind"] == "docx":
        doc = container["doc"]
        for b in blocks:
            if b["id"] not in usable:
                continue
            p = doc.paragraphs[b["para"]]
            text = XML_BAD_RE.sub("", b["result"])
            if p.runs:
                p.runs[0].text = text
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = text
        doc.save(str(out))
    else:
        parts = []
        for s in container["segs"]:
            if s["kind"] == "prose" and s["id"] in usable:
                parts.append(s["result"])
            else:
                parts.extend(s["lines"])
        out.write_text(container["newline"].join(parts), encoding="utf-8", newline="")


def default_output(path: Path) -> Path:
    return path.with_name(f"{path.stem}.lema{path.suffix}")


def write_report(path: Path, meta: dict, blocks: list[dict]) -> None:
    rows = []
    for b in blocks:
        rows.append({
            "block": b["id"], "line": b["line"], "status": b.get("status"), "attempts": b.get("attempts"),
            "provider": b.get("provider"), "missing_hard": b.get("missing_hard", []),
            "missing_soft": b.get("missing_soft", []), "errors": b.get("errors", []),
            "formatting_removed": b.get("had_formatting", False),
            "words_in": len(b["text"].split()), "words_out": len((b.get("result") or "").split()),
            "original": b["text"], "humanized": b.get("result"),
        })
    path.write_text(json.dumps({**meta, "blocks": rows}, indent=2, ensure_ascii=False), encoding="utf-8")


def print_summary(meta: dict, blocks: list[dict], kept: int) -> None:
    clean = [b for b in blocks if b.get("status") == "clean"]
    flagged = [b for b in blocks if b.get("status") == "flagged"]
    errors = [b for b in blocks if b.get("status") == "error"]
    w_in = sum(len(b["text"].split()) for b in blocks)
    w_out = sum(len((b.get("result") or "").split()) for b in blocks if b.get("status") != "error")
    print(f"Lema humanizer: {meta['engine']}")
    print(f"Input:   {meta['input']} ({len(blocks)} paragraphs rewritten, {kept} blocks left as-is)")
    print(f"Output:  {meta['output']}")
    print(f"Report:  {meta['report']}")
    took = f"  ({meta['seconds']:.0f}s)" if meta.get("seconds") else ""
    print(f"Words in rewritten paragraphs: {w_in} -> {w_out}{took}")
    print(f"Clean: {len(clean)}   Flagged: {len(flagged)}   Errors (original kept): {len(errors)}")
    for b in sorted(flagged, key=lambda b: b["line"]):
        issues = b.get("missing_hard", []) + b.get("missing_soft", [])
        print(f"  block {b['id']} (line {b['line']}): " + "; ".join(issues))
    for b in sorted(errors, key=lambda b: b["line"]):
        print(f"  block {b['id']} (line {b['line']}): ERROR {(b.get('errors') or ['unknown'])[-1]}")


# ── Agent mode: prepare / nmt / assemble ──────────────────────────────────────

STAGE_HEADERS = {
    "originals.md": "Paragraphs to rewrite (read-only). Formatting such as **bold** was removed.",
    "zh.md": "Step 1: replace each TODO with a Chinese rewrite of that block (references/agent-steps.md).",
    "ja.md": "Step 2: replace each TODO with a Japanese translation of your Chinese (references/agent-steps.md).",
}


def read_stage(path: Path) -> dict[int, str]:
    if not path.exists():
        return {}
    parts = MARKER_RE.split(path.read_text(encoding="utf-8-sig"))
    out: dict[int, str] = {}
    for k in range(1, len(parts) - 1, 2):
        body = re.sub(r"<!--.*?-->", "", parts[k + 1], flags=re.S).strip()
        if body and body != PLACEHOLDER:
            out[int(parts[k])] = " ".join(body.split())
    return out


def note_safe(s: str) -> str:
    return s.replace("--", "-").replace(">", "›")


def load_work(workdir: str) -> tuple[Path, dict]:
    work = Path(workdir).resolve()
    meta_path = work / "work.json"
    if not meta_path.exists():
        sys.exit(f"{work} is not a Lema work folder (no work.json). Run prepare first.")
    return work, json.loads(meta_path.read_text(encoding="utf-8"))


def parse_ids(spec: str | None) -> set[int] | None:
    if not spec:
        return None
    return {int(x) for x in re.split(r"[,\s]+", spec.strip()) if x}


def cmd_prepare(args) -> int:
    src = Path(args.file).resolve()
    parsed = load_source(src)
    blocks = parsed["blocks"]
    if not blocks:
        sys.exit("No prose paragraphs found to rewrite.")
    work = Path(args.workdir).resolve() if args.workdir else src.with_name(src.stem + ".lema-work")
    if work.exists():
        if not (work / "work.json").exists():
            sys.exit(f"{work} exists and is not a Lema work folder; choose another --workdir.")
        if not args.force:
            sys.exit(f"{work} already exists (it may hold drafts). Use --force to start over.")
        shutil.rmtree(work)
    work.mkdir(parents=True)
    meta = {
        "source": str(src), "created": datetime.now().isoformat(timespec="seconds"),
        "blocks": [{"id": b["id"], "line": b["line"], "original": b["text"],
                    "formatting_removed": b.get("had_formatting", False)} for b in blocks],
    }
    (work / "work.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    for name, header in STAGE_HEADERS.items():
        lines = [f"# {header}", "# Keep every <!-- block N --> marker exactly as it is.", ""]
        for b in blocks:
            lines += [f"<!-- block {b['id']} | line {b['line']} -->",
                      b["text"] if name == "originals.md" else PLACEHOLDER, ""]
        (work / name).write_text("\n".join(lines), encoding="utf-8")
    words = sum(len(b["text"].split()) for b in blocks)
    print(f"Prepared {len(blocks)} paragraphs ({words} words) from {src.name}; left as-is: " +
          ", ".join(f"{k} {v}" for k, v in sorted(parsed["kinds"].items()) if k != "prose"))
    print(f"Work folder: {work}")
    print("Next: fill zh.md, then ja.md (one block per marker), then run: nmt \"%s\"" % work)
    return 0


def cmd_nmt(args) -> int:
    work, meta = load_work(args.workdir)
    blocks = meta["blocks"]
    only = parse_ids(args.only)
    targets = [b for b in blocks if only is None or b["id"] in only]
    en_path = work / "en.md"
    if en_path.exists() and only is None and not args.force:
        sys.exit("en.md already exists (it may hold your fixes). Use --only N,M to redo specific blocks, "
                 "or --force to redo every block.")
    hops = [h for h in re.split(r"[,\s>]+", args.hops.strip()) if h]
    if len(hops) < 2:
        sys.exit("--hops needs at least two language codes, e.g. ja,fi,en")
    if hops[-1] != "en":
        log(f"warning: --hops ends in '{hops[-1]}', so the output will not be English")
    zh, ja = read_stage(work / "zh.md"), read_stage(work / "ja.md")
    problems = []
    for b in targets:
        i = b["id"]
        if i not in zh or not HAN_RE.search(zh[i]):
            problems.append(f"block {i}: zh.md has no Chinese text yet")
        if i not in ja:
            problems.append(f"block {i}: ja.md has no text yet")
        elif hops[0] == "ja" and not KANA_RE.search(ja[i]):
            problems.append(f"block {i}: ja.md text has no kana, so it is not Japanese")
    if problems:
        print("Not ready for translation:\n  " + "\n  ".join(problems))
        return 1
    for b in targets:
        lost = lost_items(protected_items(b["original"]), ja[b["id"]])
        if lost:
            log(f"warning: block {b['id']} Japanese is missing {', '.join(lost[:4])} - keep these in Latin script")

    pipe = import_pipeline()

    def one(b: dict) -> str:
        return tidy_output(pipe.heuristic_cleanup(pipe.translate_chain(ja[b["id"]], hops)))

    results: dict[int, dict] = {}
    chain = " > ".join(h.upper() for h in hops)
    log(f"Google Translate {chain} + heuristic cleanup on {len(targets)} blocks...")
    sink = sys.__stderr__ if args.verbose else open(os.devnull, "w", encoding="utf-8")
    with contextlib.redirect_stdout(sink):
        with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
            futures = {pool.submit(one, b): b for b in targets}
            for fut in as_completed(futures):
                b = futures[fut]
                try:
                    results[b["id"]] = {"text": fut.result()}
                except Exception as e:
                    results[b["id"]] = {"error": f"{type(e).__name__}: {str(e)[:200]}"}
    if not args.verbose:
        sink.close()

    existing = read_stage(en_path) if only is not None else {}
    lines = [f"# Step 3 output: after Google Translate ({chain}) and the heuristic cleanup.",
             "# Fix only what a FLAGGED note lists, keep the markers, then run assemble.", ""]
    counts = {"clean": 0, "flagged": 0, "error": 0}
    report = []
    for b in blocks:
        r = results.get(b["id"]) or ({"text": existing[b["id"]]} if b["id"] in existing else
                                     {"error": "not translated yet"})
        if "error" in r:
            status, issues, body = "error", [r["error"]], PLACEHOLDER
        else:
            hard, soft = check_integrity(b["original"], r["text"])
            issues, body = hard + soft, r["text"]
            status = "flagged" if issues else "clean"
        counts[status] += 1
        note = "clean" if status == "clean" else f"{status.upper()}: " + "; ".join(issues)
        lines += [f"<!-- block {b['id']} | line {b['line']} | {note_safe(note)} -->", body, ""]
        if status != "clean":
            report.append(f"  block {b['id']} (line {b['line']}): {status} - " + "; ".join(issues))
    en_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"{len(blocks)} blocks in en.md: {counts['clean']} clean, {counts['flagged']} flagged, "
          f"{counts['error']} errors")
    print(f"en.md: {en_path}")
    for row in report:
        print(row)
    if counts["error"]:
        print("Errors usually mean Google Translate rate-limited this IP; wait a minute and re-run "
              "with --only for those blocks.")
    print(f"Next: fix flagged blocks in en.md, then run: assemble \"{work}\"")
    return 0


def cmd_assemble(args) -> int:
    work, meta = load_work(args.workdir)
    src = Path(meta["source"])
    parsed = load_source(src)
    blocks = parsed["blocks"]
    if [b["text"] for b in blocks] != [m["original"] for m in meta["blocks"]]:
        sys.exit(f"{src.name} changed since prepare; run prepare again.")
    en = read_stage(work / "en.md")
    missing = [b["id"] for b in blocks if b["id"] not in en]
    if missing:
        sys.exit(f"en.md has no English text for block(s) {', '.join(map(str, missing))}.")
    for b in blocks:
        b["result"] = tidy_output(en[b["id"]])
        hard, soft = check_integrity(b["text"], b["result"])
        b.update(status="flagged" if hard or soft else "clean", missing_hard=hard, missing_soft=soft,
                 attempts=1, provider="agent")
    out = Path(args.output).resolve() if args.output else default_output(src)
    if out.exists() and not args.force:
        sys.exit(f"{out} already exists. Use --force to overwrite it or -o to choose another path.")
    write_document(parsed["container"], out, blocks)
    report = out.with_name(out.stem + ".report.json")
    info = {"engine": "agent mode (agent: EN > ZH > JA; Google Translate hops; heuristic cleanup)",
            "providers": ["agent", "Google Translate"], "model": None, "input": str(src),
            "output": str(out), "report": str(report), "work": str(work),
            "finished": datetime.now().isoformat(timespec="seconds")}
    write_report(report, info, blocks)
    print_summary(info, blocks, parsed["kept"])
    return 0


# ── API mode: Lema's own provider does the LLM steps ──────────────────────────

def humanize_block(block: dict, run, providers: list[str], model: str | None, retries: int) -> None:
    original = block["text"]
    best = None
    for attempt in range(1, retries + 2):
        out, used, errors = None, None, []
        for prov in providers:
            try:
                out = tidy_output(run(original, prov, model) or "")
                if out:
                    used = prov
                    break
                errors.append(f"{prov}: empty output")
            except Exception as e:  # provider failure: try the next one
                errors.append(f"{prov}: {type(e).__name__}: {str(e)[:160]}")
        if not out:
            block.setdefault("errors", []).extend(errors)
            continue
        hard, soft = check_integrity(original, out)
        score = len(hard) * 100 + len(soft)
        if best is None or score < best["score"]:
            best = {"text": out, "hard": hard, "soft": soft, "score": score, "attempt": attempt, "provider": used}
        if not hard:
            break
    if best is None:
        block.update(status="error", result=None, attempts=retries + 1)
        return
    block.update(
        status="flagged" if best["hard"] or best["soft"] else "clean",
        result=best["text"], attempts=best["attempt"], provider=best["provider"],
        missing_hard=best["hard"], missing_soft=best["soft"],
    )


def process_blocks(blocks: list[dict], run, providers, model, retries, workers, verbose) -> None:
    total = len(blocks)
    done = 0
    sink = sys.__stderr__ if verbose else open(os.devnull, "w", encoding="utf-8")
    with contextlib.redirect_stdout(sink):  # Lema's engines print progress to stdout
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            futures = {pool.submit(humanize_block, b, run, providers, model, retries): b for b in blocks}
            for fut in as_completed(futures):
                b = futures[fut]
                try:
                    fut.result()
                except Exception as e:
                    b.update(status="error", result=None, errors=[f"{type(e).__name__}: {e}"])
                done += 1
                issues = b.get("missing_hard", []) + b.get("missing_soft", [])
                note = b["status"] + (" - " + "; ".join(issues[:3]) if issues else "")
                log(f"[{done}/{total}] line {b['line']}: {note}")
    if not verbose:
        sink.close()


def cmd_api(args) -> int:
    if bool(args.file) == bool(args.text):
        sys.exit("api: give either a file or --text")
    lema_home = find_lema_home(args.lema_home)
    providers = provider_chain(args.provider, load_keys(lema_home, args.env_file), not args.no_fallback)
    run = load_engine(args.engine, lema_home, args.mode, args.intensity, args.strip_dashes)

    if args.text:
        block = {"id": 1, "text": " ".join(args.text.split()), "line": 1}
        process_blocks([block], run, providers, args.model, args.retries, 1, args.verbose)
        print(block.get("result") or block["text"])
        issues = block.get("missing_hard", []) + block.get("missing_soft", [])
        log(f"status: {block.get('status')}" + (" - " + "; ".join(issues) if issues else ""))
        return 0 if block.get("status") != "error" else 1

    src = Path(args.file).resolve()
    out = Path(args.output).resolve() if args.output else default_output(src)
    if out.exists() and not args.force:
        sys.exit(f"{out} already exists (it may contain manual fixes). Use --force or -o to choose another path.")
    parsed = load_source(src)
    blocks = parsed["blocks"]
    if not blocks:
        sys.exit("No prose paragraphs found to rewrite.")
    t0 = time.time()
    log(f"Rewriting {len(blocks)} paragraphs with Lema {args.engine} ({' > '.join(providers)}), "
        f"{args.workers} at a time...")
    process_blocks(blocks, run, providers, args.model, args.retries, args.workers, args.verbose)
    write_document(parsed["container"], out, blocks)
    report = out.with_name(out.stem + ".report.json")
    info = {"engine": f"api mode, engine {args.engine}, providers {' > '.join(providers)}",
            "providers": providers, "model": args.model, "input": str(src), "output": str(out),
            "report": str(report), "finished": datetime.now().isoformat(timespec="seconds"),
            "seconds": time.time() - t0}
    write_report(report, info, blocks)
    print_summary(info, blocks, parsed["kept"])
    return 0


# ── Other commands ────────────────────────────────────────────────────────────

def cmd_preview(args) -> int:
    parsed = load_source(Path(args.file).resolve())
    blocks = parsed["blocks"]
    print(f"{len(blocks)} paragraphs would be rewritten; left as-is: " +
          ", ".join(f"{k} {v}" for k, v in sorted(parsed["kinds"].items()) if k != "prose"))
    for b in blocks:
        words = b["text"].split()
        print(f"  block {b['id']} (line {b['line']}, {len(words)} words): {' '.join(words[:10])} ...")
    return 0


def cmd_verify(args) -> int:
    new, old = Path(args.output).resolve(), Path(args.original).resolve()
    pairs = []
    if new.suffix.lower() == ".docx":
        import docx
        old_doc, new_doc = docx.Document(str(old)), docx.Document(str(new))
        if len(old_doc.paragraphs) != len(new_doc.paragraphs):
            sys.exit("Paragraph count differs between the files; cannot line them up.")
        for t in docx_targets(old_doc):
            if t["kind"] == "prose":
                pairs.append((t["line"], t["text"], new_doc.paragraphs[t["para"]].text))
    else:
        old_segs = split_markdown(old.read_bytes().decode("utf-8-sig").replace("\r\n", "\n"))
        new_segs = split_markdown(new.read_bytes().decode("utf-8-sig").replace("\r\n", "\n"))
        classify_markdown(old_segs)
        old_blocks = [s for s in old_segs if s["kind"] != "blank"]
        new_blocks = [s for s in new_segs if s["kind"] != "blank"]
        if len(old_blocks) != len(new_blocks):
            sys.exit(f"Block count differs ({len(old_blocks)} vs {len(new_blocks)}); cannot line them up.")
        for o, n in zip(old_blocks, new_blocks):
            if o["kind"] == "prose":
                pairs.append((n["line"], o["text"], strip_inline_md(" ".join(l.strip() for l in n["lines"]))))
    bad = 0
    for line, orig, rewritten in pairs:
        hard, soft = check_integrity(orig, rewritten)
        if hard or soft:
            bad += 1
            print(f"  line {line}: " + "; ".join(hard + soft))
    print(f"Checked {len(pairs)} paragraphs: {len(pairs) - bad} clean, {bad} with issues.")
    return 0


def cmd_check_setup(args) -> int:
    ok = True
    pipe = import_pipeline()
    print("Bundled pipeline: imported (Google Translate hops and heuristic cleanup)")
    probe = "Furthermore, the findings underscores a pivotal shift—a testament to the paradigm."
    print(f"  cleanup check: {pipe.heuristic_cleanup(probe)}")
    try:
        sample = pipe.translate("Hei maailma", "fi", "en")
        print(f"Google Translate: reachable (fi 'Hei maailma' -> '{sample}')")
    except Exception as e:
        ok = False
        print(f"Google Translate: NOT reachable ({type(e).__name__}: {str(e)[:160]}).")
        print("  Agent mode needs outbound network to translate.googleapis.com.")
    try:
        import docx  # noqa: F401
        print(".docx support: python-docx installed")
    except ImportError:
        print(".docx support: python-docx missing (only .md/.txt will work)")
    try:
        lema_home = find_lema_home(args.lema_home)
    except SystemExit:
        print("API mode (optional): no Lema engine repo found, so agent mode only")
        return 0 if ok else 1
    print(f"API mode (optional): Lema engine repo at {lema_home}")
    available = load_keys(lema_home, args.env_file)
    print(f"  provider keys: {', '.join(available) or 'none'}")
    return 0 if ok else 1


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        with contextlib.suppress(Exception):
            stream.reconfigure(encoding="utf-8", errors="replace")
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--lema-home", help="api mode only: path to the Lema engine repo "
                                            "(default: LEMA_HOME or the installed pointer)")
    common.add_argument("--verbose", action="store_true", help="show the engine's own progress output")

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("prepare", parents=[common], help="split a file into blocks and create the work folder")
    p.add_argument("file")
    p.add_argument("--workdir", help="work folder (default: <name>.lema-work next to the file)")
    p.add_argument("--force", action="store_true", help="discard an existing work folder and start over")
    p.set_defaults(func=cmd_prepare)

    p = sub.add_parser("nmt", parents=[common], help="Google Translate JA > FI > EN + cleanup + checks")
    p.add_argument("workdir")
    p.add_argument("--only", help="comma-separated block numbers to redo; keeps every other block in en.md")
    p.add_argument("--hops", default="ja,fi,en",
                   help="language codes to translate through (default ja,fi,en). The first must be the "
                        "language in ja.md and the last should be en.")
    p.add_argument("--workers", type=int, default=4, help="blocks translated in parallel (default 4)")
    p.add_argument("--force", action="store_true", help="redo every block even though en.md exists")
    p.set_defaults(func=cmd_nmt)

    p = sub.add_parser("assemble", parents=[common], help="re-check en.md and write the output document")
    p.add_argument("workdir")
    p.add_argument("-o", "--output", help="output path (default: <name>.lema.<ext> next to the input)")
    p.add_argument("--force", action="store_true", help="overwrite an existing output file")
    p.set_defaults(func=cmd_assemble)

    p = sub.add_parser("api", parents=[common],
                       help="optional: let the Lema engine repo's LLM provider do everything (needs that "
                            "repo and its API keys)")
    p.add_argument("file", nargs="?", help=".md, .txt or .docx file")
    p.add_argument("--text", help="rewrite this text and print the result")
    p.add_argument("-o", "--output", help="output path (default: <name>.lema.<ext> next to the input)")
    p.add_argument("--engine", choices=["v2", "v3", "v1"], default="v2",
                   help="v2 = live-site translation chain (default); v3 = experimental LLM final hop; "
                        "v1 = actor-critic rewrite with --mode/--intensity")
    p.add_argument("--mode", default="Academic", choices=["Academic", "Professional", "Conversational", "Creative"])
    p.add_argument("--intensity", default="Enhanced", choices=["Standard", "Enhanced"])
    p.add_argument("--strip-dashes", action="store_true", help="v1 only: replace em/en dashes with commas")
    p.add_argument("--provider", default="Bedrock", choices=PROVIDER_ORDER)
    p.add_argument("--no-fallback", action="store_true", help="do not fall back to other providers on errors")
    p.add_argument("--model", help="override the provider's default model (default: gpt-oss-120b)")
    p.add_argument("--workers", type=int, default=4)
    p.add_argument("--retries", type=int, default=2,
                   help="extra attempts when a paragraph loses a citation, number, acronym or term")
    p.add_argument("--force", action="store_true", help="overwrite an existing output file")
    p.add_argument("--env-file", action="append", default=[], help="extra .env file with provider keys")
    p.set_defaults(func=cmd_api)

    p = sub.add_parser("preview", parents=[common], help="list the paragraphs that would be rewritten")
    p.add_argument("file")
    p.set_defaults(func=cmd_preview)

    p = sub.add_parser("verify", parents=[common], help="re-check an edited output against the original")
    p.add_argument("output")
    p.add_argument("original")
    p.set_defaults(func=cmd_verify)

    p = sub.add_parser("check-setup", parents=[common],
                       help="bundled pipeline, Google reachability, .docx support, optional api mode")
    p.add_argument("--env-file", action="append", default=[])
    p.set_defaults(func=cmd_check_setup)

    args = ap.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
