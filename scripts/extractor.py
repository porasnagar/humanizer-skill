"""Humanizer Multi-Modal Text Extractor.
Extracts clean, formatted text from diverse input file formats:
- Word Documents (.docx)
- PDF Documents (.pdf)
- Images with text (.png, .jpg, .jpeg, .webp)
- Plain Text & Markdown (.txt, .md)

Compatible with Antigravity, Claude Code, OpenAI Codex, OpenCode, and Hermes.
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path
from typing import Optional


def extract_from_docx(file_path: Path | str) -> str:
    """Extracts paragraphs from a .docx file, preserving heading markers and paragraph breaks."""
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required for .docx extraction. Run: pip install python-docx")

    doc = docx.Document(str(file_path))
    parts = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style_name = p.style.name.lower() if p.style else ""
        if "heading 1" in style_name:
            parts.append(f"# {txt}")
        elif "heading 2" in style_name:
            parts.append(f"## {txt}")
        elif "heading 3" in style_name:
            parts.append(f"### {txt}")
        else:
            parts.append(txt)

    # Also extract text from tables
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            parts.append("\n".join(table_rows))

    return "\n\n".join(parts)


def extract_from_pdf(file_path: Path | str) -> str:
    """Extracts clean text from PDF documents using pdfplumber or pypdf."""
    file_path = str(file_path)
    # Attempt 1: pdfplumber (best for layout and reading order)
    try:
        import pdfplumber
        pages_text = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                txt = page.extract_text()
                if txt and txt.strip():
                    pages_text.append(txt.strip())
        if pages_text:
            return "\n\n".join(pages_text)
    except Exception:
        pass

    # Attempt 2: pypdf fallback
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        pages_text = []
        for page in reader.pages:
            txt = page.extract_text()
            if txt and txt.strip():
                pages_text.append(txt.strip())
        if pages_text:
            return "\n\n".join(pages_text)
    except Exception as e:
        raise RuntimeError(f"Could not extract text from PDF {file_path}: {e}")

    return ""


def extract_from_image(file_path: Path | str) -> str:
    """Extracts text from images using pytesseract OCR, with instructions for agent vision fallback."""
    file_path = str(file_path)
    # Try pytesseract OCR if available
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
        if text and text.strip():
            return text.strip()
    except Exception:
        pass

    # Agent Vision Fallback message
    return (
        f"[IMAGE_INPUT: {file_path}]\n"
        "Note: Tesseract OCR is not installed or returned empty text. "
        "The AI Agent (Antigravity/Claude/Codex) should inspect this image directly using vision "
        "and transcribe its complete text before feeding to the engine."
    )


def extract_text(file_path: Path | str) -> str:
    """Universal text extractor routing to the appropriate format handler."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()

    if ext == ".docx":
        return extract_from_docx(path)
    elif ext == ".pdf":
        return extract_from_pdf(path)
    elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"]:
        return extract_from_image(path)
    else:
        # Default: read text/markdown with UTF-8
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError:
            with open(path, "r", encoding="latin-1") as f:
                return f.read()


def main():
    parser = argparse.ArgumentParser(description="Extract clean text from .docx, .pdf, images, or .md/.txt")
    parser.add_argument("input", type=str, help="Path to input file")
    parser.add_argument("--output", "-o", type=str, help="Optional output text file path")
    args = parser.parse_args()

    text = extract_text(args.input)
    if args.output:
        out_p = Path(args.output)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        with open(out_p, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"[Extractor] Extracted {len(text.split())} words to {args.output}")
    else:
        print(text)


if __name__ == "__main__":
    main()
